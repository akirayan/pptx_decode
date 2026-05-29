#!/usr/bin/env python3

"""
yaml2pptx.py
Build PPTX from YAML (output of: pptx2yaml.py)

Usage:
  yaml2pptx.py input.yaml output.pptx
  yaml2pptx.py input.yaml output.pptx --template template.pptx
  cat input.yaml | yaml2pptx.py - output.pptx
"""

import sys
import argparse
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))

import yaml
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import PP_PLACEHOLDER as PH

from pptx_core import create_ppt, save_presentation, set_placeholder
from pptx_components import (
    create_slide, add_blank_slide,
    add_text, add_rect, add_line, add_picture,
    BLACK, GRAY_LINE,
)


# ============================================================
# Helpers
# ============================================================

def _parse_color(hex_str):
    """'#RRGGBB' → RGBColor, None → None."""
    if not hex_str:
        return None
    h = hex_str.lstrip('#')
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _parse_align(align_str):
    return {
        'left':    PP_ALIGN.LEFT,
        'center':  PP_ALIGN.CENTER,
        'right':   PP_ALIGN.RIGHT,
        'justify': PP_ALIGN.JUSTIFY,
    }.get(align_str or 'left', PP_ALIGN.LEFT)


_PH_MAP = {
    'TITLE':        PH.TITLE,
    'BODY':         PH.BODY,
    'SUBTITLE':     PH.SUBTITLE,
    'CENTER_TITLE': PH.CENTER_TITLE,
    'OBJECT':       PH.OBJECT,
    'PICTURE':      PH.PICTURE,
}


# ============================================================
# Shape dispatcher
# ============================================================

def _dispatch(slide, s):
    fn = s.get('fn', '')

    if fn == 'add_text':
        add_text(slide,
                 s.get('text', ''),
                 s['x'], s['y'], s['w'], s['h'],
                 size=s.get('size', 12),
                 bold=s.get('bold', False),
                 color=_parse_color(s.get('color')) or BLACK,
                 align=_parse_align(s.get('align')),
                 font_name=s.get('font_name'))

    elif fn == 'add_rect':
        add_rect(slide,
                 s['x'], s['y'], s['w'], s['h'],
                 _parse_color(s.get('fill')),
                 line_color=_parse_color(s.get('line_color')),
                 line_width=s.get('line_width', 0.5))

    elif fn == 'add_line':
        add_line(slide,
                 s['x1'], s['y1'], s['x2'], s['y2'],
                 color=_parse_color(s.get('color')) or GRAY_LINE,
                 width=s.get('width', 1.0))

    elif fn == 'add_picture':
        path = s.get('image_path') or s.get('path')
        if path:
            add_picture(slide, path,
                        s['x'], s['y'],
                        w=s.get('w'), h=s.get('h'))
        else:
            print(f"  Warning: add_picture skipped — no image_path in YAML",
                  file=sys.stderr)

    elif fn == 'set_placeholder':
        ph_type = _PH_MAP.get(s.get('placeholder_type', 'TITLE'), PH.TITLE)
        try:
            set_placeholder(slide, ph_type, s.get('text', ''),
                            idx=s.get('placeholder_idx'))
        except Exception as e:
            print(f"  Warning: set_placeholder failed — {e}", file=sys.stderr)

    elif fn in ('error', 'unknown'):
        pass  # silently skip decode-time errors

    else:
        print(f"  Warning: unknown fn '{fn}' — skipped", file=sys.stderr)


# ============================================================
# Core build
# ============================================================

def build(data, template=None, output_path='output.pptx'):
    # CLI --template overrides YAML's template key
    template = template or data.get('template')
    prs = create_ppt(template)

    for slide_data in data.get('slides', []):
        layout_index = slide_data.get('layout_index', 0)
        shapes       = slide_data.get('shapes', [])

        # Slides with placeholder shapes need the layout items intact
        has_placeholder = any(s.get('fn') == 'set_placeholder' for s in shapes)
        if has_placeholder:
            slide = create_slide(prs, layout_index)
        else:
            slide = add_blank_slide(prs, layout_index)

        for s in shapes:
            try:
                _dispatch(slide, s)
            except Exception as e:
                print(f"  Warning: {s.get('fn')} ({s.get('x')},{s.get('y')}) — {e}",
                      file=sys.stderr)

    save_presentation(prs, output_path)
    #print(f"Saved: {output_path}", file=sys.stderr)


# ============================================================
# CLI
# ============================================================

def main():
    parser = argparse.ArgumentParser(
        description='Build PPTX from YAML',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Template resolution order:
  1. --template option (highest priority)
  2. 'template:' key in the YAML file
  3. No template (default blank presentation)

Examples:
  yaml2pptx.py slides.yaml output.pptx
  yaml2pptx.py slides.yaml output.pptx --template template.pptx
  cat slides.yaml | yaml2pptx.py - output.pptx
        """,
    )
    parser.add_argument('yaml_input', metavar='yaml',
                        help='YAML file path, or - for stdin')
    parser.add_argument('output',
                        help='Output PPTX file path')
    parser.add_argument('--template', default=None,
                        help='Template PPTX (overrides template key in YAML)')
    args = parser.parse_args()

    if args.yaml_input == '-':
        data = yaml.safe_load(sys.stdin)
    else:
        with open(args.yaml_input, 'r', encoding='utf-8') as f:
            data = yaml.safe_load(f)

    build(data, template=args.template, output_path=args.output)


if __name__ == '__main__':
    main()
