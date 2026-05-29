# ============================================================
# pptx_core.py
# Core primitives for PPTX manipulation (decode + build)
# ============================================================

from pptx import Presentation
from pptx.util import Cm, Pt
from pptx.enum.shapes import PP_PLACEHOLDER


# ============================================================
# 1. Presentation Layer
# ============================================================

def create_ppt(template=None):
    ppt = create_presentation(template)
    if template:
        remove_all_slides(ppt)   # keep layout/master, remove slides
    return ppt


def create_presentation(template=None):
    """Create new Presentation (with optional template)."""
    return Presentation(template) if template else Presentation()

def save_ppt(prs, path):
    """Save presentation."""
    prs.save(path)


def save_presentation(prs, path):
    """Save presentation."""
    prs.save(path)


def remove_all_slides(prs):
    """Remove all slides (keep layout/master/theme)."""
    while len(prs.slides) > 0:
        rId = prs.slides._sldIdLst[0].rId
        prs.part.drop_rel(rId)
        del prs.slides._sldIdLst[0]


# ============================================================
# 2. Layout / Slide Layer
# ============================================================

def list_layouts(prs):
    """Return list of layout names."""
    return [(i, l.name) for i, l in enumerate(prs.slide_layouts)]


def get_layout_by_name(prs, name):
    """Find layout by name."""
    for l in prs.slide_layouts:
        if l.name == name:
            return l
    raise ValueError(f"Layout not found: {name}")


def add_slide(prs, layout_index):
    """Add slide by layout index."""
    return prs.slides.add_slide(prs.slide_layouts[layout_index])


def remove_all_items(slide):
    """
    Remove all shapes from a slide (placeholders included).
    """
    spTree = slide.shapes._spTree

    # list() important (avoid mutation during iteration)
    for shape in list(slide.shapes):
        spTree.remove(shape._element)


# ============================================================
# 3. Text / Run Layer
# ============================================================

def clear_text(shape):
    """Clear shape text and return text_frame."""
    tf = shape.text_frame
    tf.clear()
    return tf


def add_paragraph(tf, level=0):
    """Add paragraph."""
    p = tf.add_paragraph()
    p.level = level
    return p


def add_run(p, text):
    """Add run."""
    run = p.add_run()
    run.text = text
    return run


def apply_paragraphs(shape, paragraphs):
    """
    Apply structured paragraphs.

    paragraphs = [
        {
            "level": 0,
            "runs": [
                {"text": "Hello"},
                {"text": "World"}
            ]
        }
    ]
    """
    tf = clear_text(shape)

    for pdata in paragraphs:
        p = add_paragraph(tf, pdata.get("level", 0))
        for rdata in pdata.get("runs", []):
            run = add_run(p, rdata.get("text", ""))


# ============================================================
# 5. Placeholder Layer (KEY)
# ============================================================

def set_placeholder(slide, ptype, text, idx=None):
    """
    Set placeholder text by (type, idx)
    type = role (TITLE, BODY, etc.)
    idx  = secondary key when multiple exist
    """
    for shape in slide.shapes:
        if not shape.is_placeholder:
            continue

        pf = shape.placeholder_format

        if pf.type == ptype:
            if idx is None or pf.idx == idx:
                shape.text = text
                return shape

    raise ValueError(f"Placeholder not found: {ptype}, idx={idx}")


# ============================================================
# 6. Decode Utilities
# ============================================================

def emu_to_cm(v):
    return round(v / 360000, 3)


def safe_get_rgb(font):
    try:
        if font.color and hasattr(font.color, "rgb") and font.color.rgb:
            return str(font.color.rgb)
    except Exception:
        pass
    return None

def safe_get_rgb(font):
    try:
        if font.color and hasattr(font.color, "rgb") and font.color.rgb:
            return f"#{str(font.color.rgb).upper()}"
    except Exception:
        pass
    return None



# ============================================================
# 7. Theme (Simplified)
# ============================================================

def get_theme_info(prs):
    """Extract minimal theme info (font only)."""
    theme_part = None

    for rel in prs.part.rels.values():
        if "theme" in rel.reltype:
            theme_part = rel._target
            break

    if theme_part is None:
        return {}

    xml = theme_part.blob.decode("utf-8")

    import re

    m = re.search(r'<a:latin typeface="([^"]+)"', xml)
    font = m.group(1) if m else None

    return {"major_font": font}


def get_effective_font(font, theme):
    """Resolve effective font (partial)."""
    return {
        "font_name": font.name if font.name else theme.get("major_font"),
        "font_size": font.size.pt if font.size else None,
        "bold": font.bold,
        "italic": font.italic,
        "color": safe_get_rgb(font),
    }


def list_theme_details(prs):
    return decode_theme(prs)


def decode_theme(prs):
    """
    Print theme details (fonts + colors).
    """
    print("=" * 60)
    print("THEME DETAILS")
    print("=" * 60)

    theme_part = None

    # find theme part
    for rel in prs.part.rels.values():
        if "theme" in rel.reltype:
            theme_part = rel._target
            break

    if theme_part is None:
        print("No theme found")
        return

    xml = theme_part.blob.decode("utf-8")

    import re

    # --- font (major / latin) ---
    major = re.search(r'<a:majorFont>.*?<a:latin typeface="([^"]+)"', xml, re.S)
    fallback = re.search(r'<a:latin typeface="([^"]+)"', xml)

    font = None
    if major:
        font = major.group(1)
    elif fallback:
        font = fallback.group(1)

    print("Major Font (Latin):", font)

    # --- east asia font ---
    east = re.search(r'<a:ea typeface="([^"]+)"', xml)
    east_font = east.group(1) if east else None

    print("East Asia Font:", east_font)

    # --- colors ---
    colors = re.findall(r'<a:srgbClr val="([^"]+)"', xml)

    print("Theme Colors (sample):", colors[:10])








# ============================================================
# 8. Decode / Debug Layer
# ============================================================


def list_shape_details(shape, theme=None, indent="  "):
    info = {
        "name": shape.name,
        "type": str(shape.shape_type),
        "x_cm": emu_to_cm(shape.left),
        "y_cm": emu_to_cm(shape.top),
        "w_cm": emu_to_cm(shape.width),
        "h_cm": emu_to_cm(shape.height),
    }

    # placeholder information
    if shape.is_placeholder:
        try:
            pf = shape.placeholder_format
            info["placeholder_type"] = str(pf.type)
            info["placeholder_idx"] = pf.idx
        except Exception:
            pass

    print(indent + str(info))

    if not shape.has_text_frame:
        return

    tf = shape.text_frame

    for pi, p in enumerate(tf.paragraphs):
        print(indent + f"  [P{pi}] text='{p.text}'")

        for ri, r in enumerate(p.runs):
            font = r.font
        
            raw = {
                "text": r.text,
                "font_name": font.name,
                "font_size": font.size.pt if font.size else None,
                "bold": font.bold,
                "italic": font.italic,
                "color": safe_get_rgb(font),
            }
        
            eff_values = get_effective_font(font, theme or {})
        
            eff = {
                "text": r.text, 
                "font_name": eff_values.get("font_name"),
                "font_size": eff_values.get("font_size"),
                "bold": eff_values.get("bold"),
                "italic": eff_values.get("italic"),
                "color": eff_values.get("color"),
            }
        
            print(indent + "    " + f"[R{ri}]")
            print(indent + "      raw:", raw)
            print(indent + "      eff:", eff)




def list_layout_details(prs):
    print("=== LAYOUT DETAILS ===")
    theme = get_theme_info(prs)
    for i, layout in enumerate(prs.slide_layouts):
        print(f"[Layout {i}] {layout.name}")
        for shape in layout.shapes:
            list_shape_details(shape, theme)


def list_slide_details(prs):
    print("=== SLIDE DETAILS ===")
    theme = get_theme_info(prs)
    for i, slide in enumerate(prs.slides):
        print(f"[Slide {i}]")
        for shape in slide.shapes:
            list_shape_details(shape, theme)



def list_master_details(prs):
    """
    Print master-level shapes (background layer).
    """
    print("=" * 60)
    print("MASTER DETAILS")
    print("=" * 60)

    theme = get_theme_info(prs)
    for i, master in enumerate(prs.slide_masters):
        print(f"\n[Master {i}]")

        for shape in master.shapes:
            list_shape_details(shape, theme)
            # print(f"  {shape.name} {shape.shape_type}")




def list_full_structure(prs):
    decode_theme(prs)
    list_master_details(prs)
    list_layout_details(prs)
    list_slide_details(prs)

