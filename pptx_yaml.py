# ============================================================
# pptx_yaml.py
# Decode PPTX slides to YAML structure for pptx_build.py
# ============================================================

import yaml
from pptx.enum.text import PP_ALIGN
from pptx_core import emu_to_cm, get_effective_font, get_theme_info


# ============================================================
# Color helpers
# ============================================================

def _rgb_str(color_obj):
    try:
        return f"#{str(color_obj.rgb).upper()}"
    except Exception:
        return None


def _extract_fill(shape):
    """Return hex fill color, or None if transparent / not solid."""
    try:
        return _rgb_str(shape.fill.fore_color)
    except Exception:
        return None


def _extract_line(shape):
    """Return {'color': hex, 'width': pt} or None if no line."""
    try:
        color = _rgb_str(shape.line.color)
        if color is None:
            return None
        width = None
        if shape.line.width is not None:
            width = round(shape.line.width.pt, 2)
        return {'color': color, 'width': width}
    except Exception:
        return None


def _align_str(align):
    if align == PP_ALIGN.CENTER:  return 'center'
    if align == PP_ALIGN.RIGHT:   return 'right'
    if align == PP_ALIGN.JUSTIFY: return 'justify'
    return 'left'


# ============================================================
# Per-shape decoders
# ============================================================

def _textbox_to_dict(shape, theme, base):
    d = {'fn': 'add_text', **base}
    tf = shape.text_frame
    if not tf.paragraphs:
        d['text'] = ''
        return d

    p = tf.paragraphs[0]
    align = _align_str(p.alignment)

    if p.runs:
        r = p.runs[0]
        eff = get_effective_font(r.font, theme)
        d['text'] = r.text
        if eff.get('font_size') is not None:
            d['size']      = eff['font_size']
        if eff.get('bold'):
            d['bold']      = True
        if eff.get('color'):
            d['color']     = eff['color']
        if eff.get('font_name'):
            d['font_name'] = eff['font_name']
    else:
        d['text'] = p.text

    if align != 'left':
        d['align'] = align

    return d


def _rect_to_dict(shape, base):
    d = {'fn': 'add_rect', **base}
    d['fill'] = _extract_fill(shape)   # None = transparent

    line = _extract_line(shape)
    if line:
        d['line_color'] = line['color']
        if line['width'] is not None:
            d['line_width'] = line['width']

    return d


def _line_to_dict(shape, base):
    """Connector bounding box → x1,y1,x2,y2."""
    d = {
        'fn': 'add_line',
        'x1': base['x'],
        'y1': base['y'],
        'x2': round(base['x'] + base['w'], 3),
        'y2': round(base['y'] + base['h'], 3),
    }
    line = _extract_line(shape)
    if line:
        d['color'] = line['color']
        if line['width'] is not None:
            d['width'] = line['width']
    return d


def _placeholder_to_dict(shape, theme, base):
    d = {'fn': 'set_placeholder', **base}
    pf = shape.placeholder_format
    d['placeholder_type'] = str(pf.type).split('(')[0].strip()
    d['placeholder_idx']  = pf.idx
    if shape.has_text_frame:
        d['text'] = shape.text_frame.text
    return d


def _shape_to_dict(shape, theme):
    base = {
        'x': emu_to_cm(shape.left),
        'y': emu_to_cm(shape.top),
        'w': emu_to_cm(shape.width),
        'h': emu_to_cm(shape.height),
    }

    if shape.is_placeholder:
        return _placeholder_to_dict(shape, theme, base)

    st = str(shape.shape_type)

    if 'TEXT_BOX'   in st: return _textbox_to_dict(shape, theme, base)
    if 'AUTO_SHAPE' in st: return _rect_to_dict(shape, base)
    if 'CONNECTOR'  in st: return _line_to_dict(shape, base)
    if 'LINE'       in st: return _line_to_dict(shape, base)
    if 'PICTURE'    in st: return {'fn': 'add_picture', **base}

    return {'fn': f'unknown', 'shape_type': st, **base}


# ============================================================
# Top-level decode
# ============================================================

def decode_to_dict(prs):
    """Decode all slides to a dict suitable for YAML serialization."""
    theme = get_theme_info(prs)

    slides = []
    for i, slide in enumerate(prs.slides):
        layout_index = None
        for j, layout in enumerate(prs.slide_layouts):
            if layout == slide.slide_layout:
                layout_index = j
                break

        shapes = []
        for shape in slide.shapes:
            try:
                shapes.append(_shape_to_dict(shape, theme))
            except Exception as e:
                shapes.append({'fn': 'error', 'message': str(e)})

        slides.append({
            'index':        i,
            'layout_index': layout_index,
            'shapes':       shapes,
        })

    return {'slides': slides}


def print_yaml(prs):
    """Print YAML representation of all slides to stdout."""
    data = decode_to_dict(prs)
    print(yaml.dump(data, allow_unicode=True, default_flow_style=False, sort_keys=False))
