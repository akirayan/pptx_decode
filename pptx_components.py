# ============================================================
# pptx_components.py
# High-level reusable PPT components
# ============================================================

from pathlib import Path
from pptx.enum.shapes import PP_PLACEHOLDER as PH
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx_core import *


# ============================================================
# Font Sizes
# ============================================================

FS_H0 = 32
FS_H1 = 24
FS_H2 = 20
FS_H3 = 16
FS_H4 = 14
FS_H5 = 12
FS_H6 = 10

# ============================================================
# Colors
# ============================================================

CS_RED_DARK  = RGBColor(0x7B, 0x00, 0x00)
CS_RED       = RGBColor(0xC0, 0x00, 0x00)
CS_RED_MID   = RGBColor(0xD3, 0x2F, 0x2F)
CS_RED_LIGHT = RGBColor(0xFF, 0xEB, 0xEE)
CS_RED_PALE  = RGBColor(0xFF, 0xF5, 0xF5)

WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
BLACK      = RGBColor(0x00, 0x00, 0x00)
GRAY_DARK  = RGBColor(0x1A, 0x1A, 0x1A)
GRAY_MID   = RGBColor(0x55, 0x55, 0x55)
GRAY_LIGHT = RGBColor(0xF5, 0xF5, 0xF5)
GRAY_LINE  = RGBColor(0xDD, 0xDD, 0xDD)

BLUE         = RGBColor(0x2D, 0x6B, 0xE4)
BLUE_LIGHT   = RGBColor(0xEE, 0xF4, 0xFF)
GREEN        = RGBColor(0x0E, 0x7C, 0x5A)
GREEN_LIGHT  = RGBColor(0xED, 0xFA, 0xF4)
ORANGE       = RGBColor(0xC0, 0x5A, 0x00)
ORANGE_LIGHT = RGBColor(0xFF, 0xF5, 0xEC)
PURPLE       = RGBColor(0x6B, 0x21, 0xA8)
PURPLE_LIGHT = RGBColor(0xF5, 0xF0, 0xFF)



# ============================================================
# 1. Basic Slide Wrappers
# ============================================================

def create_slide(prs, layout_index):
    return prs.slides.add_slide(prs.slide_layouts[layout_index])


def add_title_slide(prs, layout_index, title, subtitle=None):
    slide = create_slide(prs, layout_index)

    set_placeholder(slide, PH.TITLE, title)

    if subtitle:
        set_placeholder(slide, PH.SUBTITLE, subtitle)

    return slide


def add_content_slide(prs, layout_index, title, body=None):
    slide = create_slide(prs, layout_index)

    # Title
    set_placeholder(slide, PH.TITLE, title)

    # Body (optional)
    if body:
        set_placeholder(slide, PH.BODY, body)

    return slide


def add_section_slide(prs, layout_index, title, subtitle=None):
    slide = create_slide(prs, layout_index)

    # Section title（環境によっては TITLE / CENTER_TITLE）
    try:
        set_placeholder(slide, PH.TITLE, title)
    except Exception:
        set_placeholder(slide, PH.CENTER_TITLE, title)

    if subtitle:
        try:
            set_placeholder(slide, PH.SUBTITLE, subtitle)
        except Exception:
            pass

    return slide


def add_blank_slide(prs, layout_index):
    slide = create_slide(prs, layout_index)
    remove_all_items(slide)
    return slide





# ============================================================
# 2. Background / Header
# ============================================================

def add_full_background(slide, add_rect, color):
    add_rect(slide, 0, 0, 33.867, 19.05, fill=color)


def add_header(slide, add_rect, add_text, title, subtitle=None):
    # simplified version (you already have CS-style version)
    add_text(slide, title, 1, 1, 28, 1.5, size=20, bold=True)

    if subtitle:
        add_text(slide, subtitle, 1, 2.5, 28, 1.0, size=10)


# ============================================================
# 3. Block (core reusable unit)
# ============================================================

def add_block(slide, add_text, title, body, x, y, w, h):
    add_text(slide, title, x, y, w, 1.0, size=12, bold=True)
    add_text(slide, body, x, y + 1.2, w, h - 1.2, size=10)


# ============================================================
# 4. Bullet Panel
# ============================================================

def add_bullet_panel(slide, add_text, add_rect,
                     title, bullets,
                     x, y, w, h,
                     title_color=None, bg=None):

    add_rect(slide, x, y, w, h, fill=bg)

    add_text(slide, title, x+0.3, y+0.2, w-0.6, 1.0,
             size=11, bold=True, color=title_color)

    y_off = 1.5
    for b in bullets:
        add_text(slide, "• " + b,
                 x+0.4, y+y_off,
                 w-0.8, 0.8,
                 size=10)
        y_off += 0.8


# ============================================================
# 5. Grid Layout
# ============================================================

def add_blocks_grid(slide, blocks, add_text,
                     cols=2, x=1, y=3, w=30, h=12):

    rows = (len(blocks) + cols - 1) // cols

    cell_w = w / cols
    cell_h = h / rows

    for i, b in enumerate(blocks):
        row = i // cols
        col = i % cols

        bx = x + col * cell_w
        by = y + row * cell_h

        add_block(slide, add_text,
                  b["title"], b["body"],
                  bx, by, cell_w, cell_h)


# ============================================================
# 6. Agent Slide (IMPORTANT)
# ============================================================

def add_agent_slide(prs, layout_index,
                    add_text, add_rect,
                    title, subtitle,
                    good, risks, solutions, priorities):

    slide = create_slide(prs, layout_index)

    # background
    add_full_background(slide, add_rect, color=None)

    # header
    add_header(slide, add_rect, add_text, title, subtitle)

    # panels
    add_bullet_panel(slide, add_text, add_rect,
                     "Good", good,
                     1, 3, 10, 4)

    add_bullet_panel(slide, add_text, add_rect,
                     "Risks", risks,
                     12, 3, 20, 5)

    add_bullet_panel(slide, add_text, add_rect,
                     "Solutions", solutions,
                     1, 8, 20, 8)

    add_bullet_panel(slide, add_text, add_rect,
                     "Priority", priorities,
                     22, 8, 10, 8)

    return slide










################ others
from pptx.util import Cm, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def add_text_box(
    slide,
    text,
    x, y, w, h,
    *,
    font_size=12,
    bold=False,
    italic=False,
    color=None,
    align=PP_ALIGN.LEFT,
):
    """
    Add a simple textbox (1 paragraph / 1 run)
    """

    shape = slide.shapes.add_textbox(
        Cm(x), Cm(y), Cm(w), Cm(h)
    )

    tf = shape.text_frame
    tf.clear()

    p = tf.paragraphs[0]
    p.alignment = align

    run = p.add_run()
    run.text = text

    font = run.font
    font.size = Pt(font_size)
    font.bold = bold
    font.italic = italic

    if color:
        font.color.rgb = color

    return shape



def add_table(
    slide,
    rows,
    x, y, w, h,
    *,
    header_fill=None,
    header_text_color=None,
    body_fill=None,
    alt_fill=None,
    font_size=12,
):
    """
    Add simple table.

    rows = [
        ["Header1", "Header2"],
        ["row1 col1", "row1 col2"]
    ]
    """

    if not rows:
        raise ValueError("rows must not be empty")

    n_rows = len(rows)
    n_cols = len(rows[0])

    shape = slide.shapes.add_table(
        n_rows,
        n_cols,
        Cm(x), Cm(y), Cm(w), Cm(h)
    )

    table = shape.table

    for r in range(n_rows):
        for c in range(n_cols):

            cell = table.cell(r, c)
            text = rows[r][c]

            tf = cell.text_frame
            tf.clear()

            p = tf.paragraphs[0]
            run = p.add_run()
            run.text = text
            run.font.size = Pt(font_size)

            # -----------------
            # Header row
            # -----------------
            if r == 0:
                if header_fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = header_fill

                if header_text_color:
                    run.font.color.rgb = header_text_color

            # -----------------
            # Body rows
            # -----------------
            else:
                if alt_fill and r % 2 == 1:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = alt_fill
                elif body_fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = body_fill

    return table




def add_rect(slide, x, y, w, h, fill, line_color=None, line_width=0.5):
    shape = slide.shapes.add_shape(1, Cm(x), Cm(y), Cm(w), Cm(h))
    if fill is None:
        shape.fill.background()
    else:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill
    if line_color is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line_color
        shape.line.width = Pt(line_width)
    return shape


def add_text(slide, text, x, y, w, h, size=12, bold=False, color=BLACK,
             align=PP_ALIGN.LEFT, font_name=None):
    shape = slide.shapes.add_textbox(Cm(x), Cm(y), Cm(w), Cm(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    if font_name:
        run.font.name = font_name
    return shape


def add_picture(slide, image_path, x, y, w=None, h=None):
    return slide.shapes.add_picture(
        str(image_path),
        Cm(x), Cm(y),
        width=Cm(w) if w is not None else None,
        height=Cm(h) if h is not None else None,
    )


def add_line(slide, x1, y1, x2, y2, color=GRAY_LINE, width=1.0):
    line = slide.shapes.add_connector(1, Cm(x1), Cm(y1), Cm(x2), Cm(y2))
    line.line.color.rgb = color
    line.line.width = Pt(width)
    return line


def add_note(slide, text, y=17.75):
    return add_text(slide, text, 1.2, y, 31.0, 0.5,
                    size=8.5, color=GRAY_MID, align=PP_ALIGN.CENTER)


def add_card(slide, title, items, x, y, w, h,
             fill=GRAY_LIGHT,
             title_color=BLACK):

    add_rect(slide, x, y, w, h, fill)

    add_text(slide, title,
             x + 0.4, y + 0.3, w - 0.8, 1.0,
             size=FS_H3, bold=True, color=title_color)

    yoff = y + 1.4
    for item in items:
        add_text(slide, "• " + item,
                 x + 0.5, yoff,
                 w - 1.0, 0.8,
                 size=FS_H4, color=GRAY_DARK)
        yoff += 0.8


# ============================================================
# 8. Content Header / Bullet Block
# ============================================================

def add_content_header(slide, title, subtitle=None, font_name=None):
    add_text(slide, title,    1.50, 0.50, 26.00, 1.20, size=FS_H1, bold=True, color=BLACK, font_name=font_name)
    if subtitle:
        add_text(slide, subtitle, 1.80, 1.70, 26.00, 0.80, size=FS_H3, color=BLACK, font_name=font_name)


def add_bullet_block(slide, title, bullets, l, t, w, h,
                     title_color=CS_RED, bg=CS_RED_PALE, bullet_color=GRAY_DARK,
                     font_name=None):
    add_rect(slide, l, t, w, h, fill=bg, line_color=CS_RED_LIGHT, line_width=0.5)
    add_text(slide, title, l+0.30, t+0.23, w-0.56, 0.86,
             size=FS_H3, bold=True, color=title_color, font_name=font_name)
    y_off = 1.12
    for b in bullets:
        if y_off + 0.75 > h:
            break
        add_text(slide, "• " + b, l+0.43, t+y_off, w-0.69, 0.76,
                 size=FS_H4, color=bullet_color, font_name=font_name)
        y_off += 1.3

